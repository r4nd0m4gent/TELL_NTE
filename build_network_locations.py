"""
build_network_locations.py
──────────────────────────
Write site/network/data/org_locations.json: where the organisations in the
funding network are, taken from the CORDIS project exports (Horizon Europe and
Horizon 2020). CORDIS gives every participant a city, coordinates and a type.

    HES  higher or secondary education     PRC  private company
    REC  research organisation             PUB  public body       OTH  other

The page looks organisations up by name and country ("NAME|CC", name in
capitals with single spaces), so Erasmus+ and NWO partners are found too when
they also take part in a Horizon project under the same name. NWO lists no
countries; its partners are looked up as Dutch. Only organisations that occur
in the page's data are written.

Dutch organisations are placed by data/nl_org_locations.json first (their
main seat); this file fills in the rest, and places the European partners.

The exports (~150 MB) are downloaded to --cordis-dir once and reused.

Usage
-----
    python build_network_locations.py
    python build_network_locations.py --cordis-dir data/cordis
"""

from __future__ import annotations

import argparse
import io
import json
import urllib.request
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "site" / "network" / "data"
OUT = DATA / "org_locations.json"
EXPORTS = {   # same exports as textile-network/scrapers/cordis.py
    "cordis-HORIZONprojects-xlsx.zip": "https://cordis.europa.eu/data/cordis-HORIZONprojects-xlsx.zip",
    "cordis-h2020projects-xlsx.zip": "https://cordis.europa.eu/data/cordis-h2020projects-xlsx.zip",
}


def key(name: str | None, country: str | None) -> str:
    return " ".join(str(name or "").upper().split()) + "|" + str(country or "").upper().strip()


def wanted_keys() -> set[str]:
    """Name|country of every organisation in the page's data."""
    keys = set()
    for fname in ("cordis_textile_projects.json", "erasmus_textile_projects.json"):
        data = json.loads((DATA / fname).read_text(encoding="utf-8"))
        for p in data["projects"]:
            for m in p.get("project_members") or []:
                keys.add(key(m.get("organisation"), m.get("country")))
    for p in json.loads((DATA / "nwo_textile_raw.json").read_text(encoding="utf-8")):
        for m in p.get("project_members") or []:
            keys.add(key((m.get("organisation") or "").split("||")[0], "NL"))
    return keys


def export(cordis_dir: Path, fname: str) -> Path:
    path = cordis_dir / fname
    if not path.exists():
        cordis_dir.mkdir(parents=True, exist_ok=True)
        print(f"downloading {EXPORTS[fname]} …")
        urllib.request.urlretrieve(EXPORTS[fname], path)
    return path


def participants(zip_path: Path):
    """(name, country, city, lat, lon, type) for each row of organization.xlsx."""
    with zipfile.ZipFile(zip_path) as z:
        wb = load_workbook(io.BytesIO(z.read("organization.xlsx")), read_only=True)
    rows = wb.worksheets[0].iter_rows(values_only=True)
    header = next(rows)
    for row in rows:
        r = dict(zip(header, row))
        try:
            lat, lon = (float(v) for v in str(r.get("geolocation") or "").split(","))
        except ValueError:
            continue
        yield r.get("name"), r.get("country"), (r.get("city") or "").strip(), lat, lon, r.get("activityType") or ""


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--cordis-dir", type=Path, default=ROOT / "data" / "cordis",
                   help="where the CORDIS exports are kept (default: data/cordis)")
    args = p.parse_args()

    wanted = wanted_keys()
    # An organisation appears once per project; keep its most frequent location
    seen: dict[str, Counter] = defaultdict(Counter)
    for fname in EXPORTS:
        for name, country, city, lat, lon, typ in participants(export(args.cordis_dir, fname)):
            k = key(name, country)
            if k in wanted:
                seen[k][(city, round(lat, 4), round(lon, 4), typ)] += 1

    found = {k: list(c.most_common(1)[0][0]) for k, c in sorted(seen.items())}
    # [city, lat, lon, type] per organisation, compact
    out = {
        "_about": ("Locations of the funding network's organisations from the CORDIS exports "
                   "(Horizon Europe, Horizon 2020). Key: NAME|COUNTRY, name in capitals. "
                   "Value: [city, lat, lon, activity type (HES/REC/PRC/PUB/OTH)]. "
                   "Written by build_network_locations.py."),
        "organisations": {k: [city, lat, lon, typ] for k, (city, lat, lon, typ) in found.items()},
    }
    OUT.write_text(json.dumps(out, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    types = Counter(v[3] for v in out["organisations"].values())
    print(f"{len(found)} of {len(wanted)} organisations located -> {OUT.relative_to(ROOT)}"
          f" ({OUT.stat().st_size/1024:.0f} KB); types {dict(types)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
